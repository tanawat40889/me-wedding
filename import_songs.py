#!/usr/bin/env python3
# แปลง CSV/XLSX -> songs.js   ใช้: python3 import_songs.py songs.csv
import csv, json, sys, pathlib

src = pathlib.Path(sys.argv[1] if len(sys.argv) > 1 else "songs-template.csv")
if src.suffix.lower() == ".xlsx":
    from openpyxl import load_workbook
    rows = [[("" if c is None else str(c)) for c in r] for r in load_workbook(src, read_only=True).active.iter_rows(values_only=True)]
else:
    rows = list(csv.reader(src.open(encoding="utf-8-sig")))

GEN = {"ผู้ใหญ่", "วัยทำงาน", "เด็ก"}
LV = {"ง่าย", "กลาง", "ยาก"}
songs = []
for n, r in enumerate(rows[1:], 2):
    r = [x.strip() for x in r] + [""] * 7
    title, artist, gen, lv, lang, yt, lyrics = r[:7]
    if not title or title.startswith("ตัวอย่าง"):
        continue
    if gen not in GEN: print(f"แถว {n}: วัย '{gen}' ไม่รู้จัก → ใช้ วัยทำงาน"); gen = "วัยทำงาน"
    if lv not in LV: print(f"แถว {n}: ระดับ '{lv}' ไม่รู้จัก → ใช้ กลาง"); lv = "กลาง"
    if not lyrics: print(f"แถว {n}: '{title}' ยังไม่มีเนื้อเพลง")
    songs.append({"title": title, "artist": artist, "gen": gen, "level": lv,
                  "lang": "en-US" if lang.lower() in ("อังกฤษ", "en", "english", "สากล") else "th-TH",
                  "lyrics": "\n".join(l.strip() for l in lyrics.replace("\r", "").split("\n") if l.strip()),
                  "yt": yt})

out = pathlib.Path(__file__).with_name("songs.js")
teams = '["ทีมเจ้าบ่าว", "ทีมเจ้าสาว"]'
if out.exists():
    import re
    m = re.search(r"window\.TEAMS\s*=\s*(\[.*?\]);", out.read_text(encoding="utf-8"))
    if m: teams = m.group(1)
out.write_text("// สร้างจาก " + src.name + " ด้วย import_songs.py\nwindow.SONGS = "
               + json.dumps(songs, ensure_ascii=False, indent=1) + ";\n\n// ชื่อทีม\nwindow.TEAMS = " + teams + ";\n", encoding="utf-8")
print(f"✓ เขียน {len(songs)} เพลง → {out.name}")
